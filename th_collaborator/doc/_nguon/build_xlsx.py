# -*- coding: utf-8 -*-
"""Gom toàn bộ tài liệu vào 1 file Excel."""
import sys

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

import content as c

OUT, IMG_DIR = sys.argv[1], sys.argv[2]

FONT = 'Arial'
HEAD_FILL = PatternFill('solid', fgColor='00695C')
SECTION_FILL = PatternFill('solid', fgColor='E0F2F1')
INPUT_FILL = PatternFill('solid', fgColor='FFF9C4')
THIN = Side(style='thin', color='B0BEC5')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical='top')

wb = Workbook()
wb.remove(wb.active)


def title_block(ws, title, subtitle=None):
    ws['A1'] = title
    ws['A1'].font = Font(name=FONT, size=15, bold=True, color='00695C')
    ws['A2'] = subtitle or f'{c.PROJECT["name"]} · {c.PROJECT["modules"]} · Cập nhật {c.PROJECT["date"]}'
    ws['A2'].font = Font(name=FONT, size=9, italic=True, color='607D8B')
    return 4


def write_table(ws, row, headers, rows, widths=None, input_cols=()):
    """Ghi bảng có header; trả về (dòng đầu dữ liệu, dòng cuối)."""
    for col, head in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=head)
        cell.font = Font(name=FONT, bold=True, color='FFFFFF')
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
        cell.border = BORDER
    for r_index, values in enumerate(rows, row + 1):
        for col, value in enumerate(values, 1):
            cell = ws.cell(row=r_index, column=col, value=value)
            cell.font = Font(name=FONT, size=10)
            cell.alignment = WRAP
            cell.border = BORDER
            if col in input_cols:
                cell.fill = INPUT_FILL
    if widths:
        for col, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(col)].width = width
    return row + 1, row + len(rows)


def section(ws, row, text, span):
    ws.cell(row=row, column=1, value=text).font = Font(name=FONT, size=12, bold=True, color='004D40')
    for col in range(1, span + 1):
        ws.cell(row=row, column=col).fill = SECTION_FILL
    return row + 1


def finish(ws, header_row, first, last, ncols):
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    ws.auto_filter.ref = f'A{header_row}:{get_column_letter(ncols)}{max(last, first)}'
    ws.sheet_view.zoomScale = 100


def status_rules(ws, rng, col_letter, first):
    green = PatternFill('solid', fgColor='C8E6C9')
    amber = PatternFill('solid', fgColor='FFE0B2')
    red = PatternFill('solid', fgColor='FFCDD2')
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT(${col_letter}{first},3)="Đạt"'], fill=green))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT(${col_letter}{first},9)="Chưa test"'], fill=amber))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'${col_letter}{first}="Không đạt"'], fill=red))


# --- Sheets dữ liệu (tạo trước để Tổng quan tham chiếu) --------------------
ws_ov = wb.create_sheet('Tổng quan')

# Thiết kế
ws = wb.create_sheet('Thiết kế')
row = title_block(ws, '01. THIẾT KẾ HỆ THỐNG')
row = section(ws, row, 'Kiến trúc 2 module', 4)
_, last = write_table(ws, row, ['Module', 'Vai trò', 'Phụ thuộc', 'Ghi chú'], [
    ('th_collaborator', 'Lõi nghiệp vụ + khung ký số độc lập nhà cung cấp', 'mail', 'Version 16.0.1.3.0'),
    ('th_collaborator_sign_vnpt', 'Nhà cung cấp ký số VNPT-CA Plugin (USB Token)', 'th_collaborator', 'Version 16.0.1.0.0'),
], widths=[34, 40, 46, 70])
row = section(ws, last + 2, 'Thành phần', 4)
_, last = write_table(ws, row, ['Thành phần', 'Loại', 'File chính', 'Trách nhiệm'], c.COMPONENTS)
row = section(ws, last + 2, 'Giao tiếp HTTP', 4)
_, last = write_table(ws, row, ['Route', 'Auth', 'Method', 'Mô tả'], c.ROUTES)
row = section(ws, last + 2, 'Bảo mật & phân quyền', 4)
_, last = write_table(ws, row, ['Hạng mục', 'Thiết kế', 'Ghi chú', ''], [s + ('',) for s in c.SECURITY])
row = section(ws, last + 2, 'Mở rộng nhà cung cấp ký số', 4)
_, last = write_table(ws, row, ['Bước', 'Nội dung', '', ''], [(i, s, '', '') for i, s in enumerate(c.EXTENSION_STEPS, 1)])
for r in range(row + 1, last + 1):
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
ws.freeze_panes = 'A4'

# Mô hình dữ liệu
ws = wb.create_sheet('Mô hình dữ liệu')
row = title_block(ws, '01b. MÔ HÌNH DỮ LIỆU')
rows = [(model, desc, f, t, m) for model, desc, fields in c.MODELS for f, t, m in fields]
first, last = write_table(ws, row, ['Model', 'Mô tả model', 'Trường', 'Kiểu', 'Ý nghĩa'], rows, widths=[36, 34, 44, 30, 60])
finish(ws, row, first, last, 5)

# Sơ đồ
ws = wb.create_sheet('Sơ đồ')
title_block(ws, 'SƠ ĐỒ TỔNG QUAN & QUY TRÌNH')
for name, anchor, width in (('so_do_tong_quan.png', 'A4', 1100), ('so_do_quy_trinh.png', 'A62', 1100)):
    image = XLImage(f'{IMG_DIR}/{name}')
    ratio = width / image.width
    image.width, image.height = width, int(image.height * ratio)
    ws.add_image(image, anchor)

# Yêu cầu nghiệp vụ
ws = wb.create_sheet('YC nghiệp vụ')
row = title_block(ws, '02a. YÊU CẦU NGHIỆP VỤ (BR)')
first, last = write_table(ws, row, ['ID', 'Yêu cầu', 'Mục tiêu / lý do', 'Ưu tiên', 'Chức năng đáp ứng'],
                          [br + (', '.join(f[0] for f in c.FUNC_REQS if br[0] in f[3]) or 'NFR-05',) for br in c.BUSINESS_REQS],
                          widths=[9, 70, 40, 12, 30])
finish(ws, row, first, last, 5)
BR_RANGE = (first, last)

# Yêu cầu chức năng
ws = wb.create_sheet('YC chức năng')
row = title_block(ws, '02b. YÊU CẦU CHỨC NĂNG (FR)')
first, last = write_table(ws, row, ['ID', 'Chức năng', 'Mô tả / tiêu chí chấp nhận', 'BR', 'Trạng thái'], c.FUNC_REQS,
                          widths=[9, 28, 90, 14, 26], input_cols=(5,))
finish(ws, row, first, last, 5)
FR_RANGE = (first, last)
dv = DataValidation(type='list', formula1='"Hoàn thành,Hoàn thành (chưa test token thật),Đang làm,Chưa làm"', allow_blank=True)
ws.add_data_validation(dv)
dv.add(f'E{first}:E{last + 30}')

# Yêu cầu phi chức năng
ws = wb.create_sheet('YC phi chức năng')
row = title_block(ws, '02c. YÊU CẦU PHI CHỨC NĂNG (NFR)')
first, last = write_table(ws, row, ['ID', 'Nhóm', 'Yêu cầu'], c.NON_FUNC_REQS, widths=[9, 22, 110])
finish(ws, row, first, last, 3)

# Hướng dẫn sử dụng
ws = wb.create_sheet('Hướng dẫn sử dụng')
row = title_block(ws, '03a. HƯỚNG DẪN SỬ DỤNG')
rows = [(f'A{i}. {sec}', j, step) for i, (sec, steps) in enumerate(c.USER_GUIDE, 1) for j, step in enumerate(steps, 1)]
first, last = write_table(ws, row, ['Mục', 'Bước', 'Thao tác'], rows, widths=[34, 7, 110])
finish(ws, row, first, last, 3)

# Kịch bản kiểm thử
ws = wb.create_sheet('Kịch bản kiểm thử')
row = title_block(ws, '03b. KỊCH BẢN KIỂM THỬ',
                  'Ô nền vàng (Trạng thái, Người test, Ngày test, Ghi chú) để đội test cập nhật. Môi trường đã test: DB th_collaborator_test, plugin VNPT giả lập.')
rows = [tc + ('Claude (DB test)' if tc[6].startswith('Đạt') else '', '16-17/10/2026' if tc[6].startswith('Đạt') else '', '') for tc in c.TEST_CASES]
first, last = write_table(ws, row, ['ID', 'FR', 'Điều kiện', 'Các bước', 'Kết quả mong đợi', 'Kết quả thực tế', 'Trạng thái',
                                    'Người test', 'Ngày test', 'Ghi chú'], rows,
                          widths=[8, 10, 26, 42, 42, 34, 20, 16, 13, 24], input_cols=(7, 8, 9, 10))
finish(ws, row, first, last, 10)
TC_RANGE = (first, last)
dv = DataValidation(type='list', formula1='"Đạt,Đạt (plugin giả),Không đạt,Chưa test,Chưa test – cần token"', allow_blank=True)
ws.add_data_validation(dv)
dv.add(f'G{first}:G{last + 50}')
status_rules(ws, f'A{first}:J{last + 50}', 'G', first)

# Bảo trì
ws = wb.create_sheet('Bảo trì')
row = title_block(ws, '04a. HƯỚNG DẪN BẢO TRÌ')
first, last = write_table(ws, row, ['Hạng mục', 'Nội dung'], c.MAINTENANCE, widths=[24, 120])
finish(ws, row, first, last, 2)

# Tồn đọng
ws = wb.create_sheet('Tồn đọng')
row = title_block(ws, '04b. VẤN ĐỀ TỒN ĐỌNG', 'Ô nền vàng (Trạng thái xử lý, Hạn, Ghi chú) để cập nhật tiến độ.')
rows = [iss + ('Mở', '', '') for iss in c.OPEN_ISSUES]
first, last = write_table(ws, row, ['ID', 'Vấn đề', 'Mức độ', 'Ảnh hưởng', 'Đề xuất', 'Phụ trách', 'Trạng thái xử lý', 'Hạn', 'Ghi chú'],
                          rows, widths=[9, 52, 11, 34, 46, 16, 16, 12, 22], input_cols=(7, 8, 9))
finish(ws, row, first, last, 9)
ISS_RANGE = (first, last)
dv = DataValidation(type='list', formula1='"Mở,Đang xử lý,Đã xử lý,Không xử lý"', allow_blank=True)
ws.add_data_validation(dv)
dv.add(f'G{first}:G{last + 50}')
dv_sev = DataValidation(type='list', formula1='"Cao,Trung bình,Thấp"', allow_blank=True)
ws.add_data_validation(dv_sev)
dv_sev.add(f'C{first}:C{last + 50}')
ws.conditional_formatting.add(f'C{first}:C{last + 50}', FormulaRule(formula=[f'$C{first}="Cao"'], fill=PatternFill('solid', fgColor='FFCDD2')))
ws.conditional_formatting.add(f'C{first}:C{last + 50}', FormulaRule(formula=[f'$C{first}="Trung bình"'], fill=PatternFill('solid', fgColor='FFE0B2')))
ws.conditional_formatting.add(f'G{first}:G{last + 50}', FormulaRule(formula=[f'$G{first}="Đã xử lý"'], fill=PatternFill('solid', fgColor='C8E6C9')))

# --- Tổng quan (công thức tham chiếu các sheet) ------------------------------
ws = ws_ov
title_block(ws, 'TỔNG QUAN DỰ ÁN TH COLLABORATOR')
info = [('Dự án', c.PROJECT['name']), ('Module', c.PROJECT['modules']), ('Nền tảng', c.PROJECT['platform']),
        ('Thư mục code', c.PROJECT['path']), ('Tác giả', c.PROJECT['author']), ('Cập nhật', c.PROJECT['date'])]
row = section(ws, 4, 'Thông tin chung', 3)
write_table(ws, row, ['Mục', 'Giá trị', ''], [i + ('',) for i in info], widths=[34, 70, 30])
row = section(ws, row + len(info) + 2, 'Số liệu (tự cập nhật theo các sheet)', 3)
fr0, fr1 = FR_RANGE
tc0, tc1 = TC_RANGE
is0, is1 = ISS_RANGE
br0, br1 = BR_RANGE
stats = [
    ('Yêu cầu nghiệp vụ', f"=COUNTA('YC nghiệp vụ'!A{br0}:A{br1 + 50})", 'Số dòng BR'),
    ('Yêu cầu chức năng', f"=COUNTA('YC chức năng'!A{fr0}:A{fr1 + 30})", 'Số dòng FR'),
    ('  - Hoàn thành', f"=COUNTIF('YC chức năng'!E{fr0}:E{fr1 + 30},\"Hoàn thành*\")", 'Trạng thái bắt đầu "Hoàn thành"'),
    ('Kịch bản kiểm thử', f"=COUNTA('Kịch bản kiểm thử'!A{tc0}:A{tc1 + 50})", ''),
    ('  - Đạt', f"=COUNTIF('Kịch bản kiểm thử'!G{tc0}:G{tc1 + 50},\"Đạt*\")", 'Gồm "Đạt (plugin giả)"'),
    ('  - Không đạt', f"=COUNTIF('Kịch bản kiểm thử'!G{tc0}:G{tc1 + 50},\"Không đạt\")", ''),
    ('  - Chưa test', f"=COUNTIF('Kịch bản kiểm thử'!G{tc0}:G{tc1 + 50},\"Chưa test*\")", ''),
    ('  - Tỉ lệ đạt', None, 'Đạt / tổng kịch bản'),
    ('Vấn đề tồn đọng', f"=COUNTA('Tồn đọng'!A{is0}:A{is1 + 50})", ''),
    ('  - Mức Cao', f"=COUNTIF('Tồn đọng'!C{is0}:C{is1 + 50},\"Cao\")", ''),
    ('  - Mức Trung bình', f"=COUNTIF('Tồn đọng'!C{is0}:C{is1 + 50},\"Trung bình\")", ''),
    ('  - Mức Thấp', f"=COUNTIF('Tồn đọng'!C{is0}:C{is1 + 50},\"Thấp\")", ''),
    ('  - Còn mở', f"=COUNTIF('Tồn đọng'!G{is0}:G{is1 + 50},\"Mở\")+COUNTIF('Tồn đọng'!G{is0}:G{is1 + 50},\"Đang xử lý\")", 'Mở + Đang xử lý'),
]
first, last = write_table(ws, row, ['Chỉ số', 'Giá trị', 'Cách tính'], stats)
total_row = first + 3
pass_row = first + 4
rate_row = first + 7
ws.cell(row=rate_row, column=2, value=f'=IF(B{total_row}=0,0,B{pass_row}/B{total_row})').number_format = '0.0%'
for r in range(first, last + 1):
    ws.cell(row=r, column=2).alignment = Alignment(horizontal='left', vertical='top')
    ws.cell(row=r, column=2).font = Font(name=FONT, size=10, bold=True)
row = section(ws, last + 2, 'Danh mục sheet', 3)
sheets = [('Thiết kế', 'Kiến trúc, thành phần, route, bảo mật, mở rộng ký số'),
          ('Mô hình dữ liệu', 'Model & trường'), ('Sơ đồ', 'Sơ đồ tổng quan + quy trình'),
          ('YC nghiệp vụ', 'BR + truy vết FR'), ('YC chức năng', 'FR + trạng thái'), ('YC phi chức năng', 'NFR'),
          ('Hướng dẫn sử dụng', 'Thao tác theo chức năng'), ('Kịch bản kiểm thử', 'TC + kết quả (cập nhật được)'),
          ('Bảo trì', 'Vận hành, cấu hình, xử lý lỗi'), ('Tồn đọng', 'Vấn đề + đề xuất (cập nhật được)')]
first_s, last_s = write_table(ws, row, ['Sheet', 'Nội dung', ''], [s + ('',) for s in sheets])
for r, (name, _) in zip(range(first_s, last_s + 1), sheets):
    cell = ws.cell(row=r, column=1)
    cell.hyperlink = f"#'{name}'!A1"
    cell.font = Font(name=FONT, size=10, color='1565C0', underline='single')
row = section(ws, last_s + 2, 'Các giai đoạn đã thực hiện', 3)
first_t, last_t = write_table(ws, row, ['Giai đoạn', 'Nội dung', ''], [t + ('',) for t in c.TIMELINE])
row = section(ws, last_t + 2, 'Chú thích', 3)
note = ws.cell(row=row, column=1, value='Ô nền vàng')
note.fill = INPUT_FILL
note.font = Font(name=FONT, size=10)
ws.cell(row=row, column=2, value='Cột cho phép cập nhật (trạng thái, người test, hạn xử lý...). Các ô khác là nội dung tài liệu.').font = Font(name=FONT, size=10)
ws.freeze_panes = 'A4'

# Excel tự tính công thức khi mở (không có LibreOffice để tính sẵn)
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print('saved', OUT, wb.sheetnames)
